#!/usr/bin/env python3

import argparse
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

def hex_to_rgb(hex_str: str) -> RGBColor:
    s = (hex_str or "").strip().lstrip("#")
    if len(s) == 3:
        s = "".join(ch * 2 for ch in s)
    try:
        r = int(s[0:2], 16)
        g = int(s[2:4], 16)
        b = int(s[4:6], 16)
        return RGBColor(r, g, b)
    except Exception:
        return RGBColor(47, 85, 151)  # default accent

def add_accent_bar(slide, color: RGBColor):
    try:
        bar = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.25)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        if bar.line:
            bar.line.fill.background()
    except Exception:
        pass

def style_title_shape(shape, color: RGBColor):
    if not shape or not hasattr(shape, "text_frame"):
        return
    tf = shape.text_frame
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(44)
            r.font.bold = True
            r.font.color.rgb = color

def style_body_frame(tf):
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(20)
            r.font.color.rgb = RGBColor(60, 60, 60)

def add_footer(slide, text: str):
    if not text:
        return
    left = Inches(0.5)
    top = Inches(7.0)
    width = Inches(12.33)
    height = Inches(0.4)
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.RIGHT
    for r in p.runs:
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(100, 100, 100)


def add_title_slide(prs: Presentation, title: str, subtitle: str):
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    return slide


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str], level: int = 0):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    tx_box = slide.placeholders[1]
    tf = tx_box.text_frame
    tf.clear()

    first = True
    for item in bullets:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.text = item
        p.level = level
    # style bullets
    style_body_frame(tf)
    return slide


def add_title_only_slide(prs: Presentation, title: str, lines: list[str] | None = None):
    slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title

    if lines:
        left = Inches(1)
        top = Inches(2.0)
        width = Inches(8)
        height = Inches(3)
        tx_box = slide.shapes.add_textbox(left, top, width, height)
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = line
            p.level = 0


def build_deck(args):
    project_name = args.project_name or "Room Matcher AI"
    one_liner = args.tagline or (
        "No Hustle for Hostels - Roommate AI Karo Na"
    )
    author = args.author or ""

    prs = Presentation()
    accent_rgb = hex_to_rgb(getattr(args, "accent", "#2F5597"))
    footer_text = args.footer or args.website or project_name

    subtitle_lines = [one_liner]
    if author:
        subtitle_lines.append(f"By {author}")
    subtitle_lines.append(datetime.now().strftime("%b %d, %Y"))
    title_slide = add_title_slide(prs, f"{project_name}", "\n".join(subtitle_lines))
    style_title_shape(title_slide.shapes.title, accent_rgb)
    add_accent_bar(title_slide, accent_rgb)
    # optional logo
    if getattr(args, "logo", ""):
        lp = Path(args.logo)
        if lp.exists():
            title_slide.shapes.add_picture(str(lp), Inches(11.8), Inches(0.35), width=Inches(1.2))
    add_footer(title_slide, footer_text)

    s = add_bullets_slide(
        prs,
        "Problem",
        [
            "Finding compatible roommates is time-consuming and risky",
            "Unstructured ads and scattered data make matching hard",
            "Budget/city constraints + lifestyle preferences complicate discovery",
            "Limited explainability and scam risk in current flows",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    s = add_bullets_slide(
        prs,
        "Solution",
        [
            "A FastAPI backend that parses free text into structured profiles",
            "Multi-agent pipeline to score compatibility and explain matches",
            "Search rooms via free-text or structured filters with fuzzy amenities",
            "Optional degraded mode for offline/low-bandwidth environments",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    s = add_bullets_slide(
        prs,
        "Product Highlights",
        [
            "Explainability with reasons, flags (noise/smoking/pets/budget), suggestions",
            "City & budget-aware room suggestions, quantile defaults",
            "Fuzzy amenity mapping and weighting (AC, wifi, Parking, etc.)",
            "CORS-enabled for local frontend integration",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    s = add_bullets_slide(
        prs,
        "How It Works (Pipeline)",
        [
            "Profile Reader: parse city, budget, cleanliness, sleep, noise, smoking/pets, food",
            "Match Scorer: composite of sleep/clean/noise/budget/special",
            "Red Flag Agent: detects potential conflicts + suspicious content",
            "Wingman: explanations + suggestions",
            "Room Hunter: city/budget intersect, rent + amenities scoring",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    s = add_bullets_slide(
        prs,
        "Degraded Mode (Low-Compute)",
        [
            "Prefilters by city + budget IoU, caps candidate pool",
            "Simplified scoring (budget + cleanliness)",
            "Minimal flags (budget mismatch)",
            "Optional room suggestions (configurable)",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    s = add_bullets_slide(
        prs,
        "Target Users & Use Cases",
        [
            "Students & young professionals in major Pakistani cities",
            "Hostels/PGs and property managers seeking qualified leads",
            "Portals integrating AI matching for conversion uplift",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    s = add_bullets_slide(
        prs,
        "Tech & Architecture",
        [
            "FastAPI, Uvicorn, Pydantic",
            "Dataset: synthetic profiles and listings JSON",
            "CORS for local frontend (e.g., Next.js @ http://localhost:3000)",
            "Dockerized deployment; Makefile automation",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    s = add_bullets_slide(
        prs,
        "Data & Metrics",
        [
            "Datasets: 400 roommate profiles, 400 listings (synthetic)",
            "Match score components: sleep, clean, noise, budget, special",
            "Room search scores: rent proximity + amenity match",
            "Quality levers: amenity weights, min score cutoff",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    s = add_bullets_slide(
        prs,
        "Business Model & GTM (Draft)",
        [
            "API access for portals (SaaS, per-call or tiered)",
            "Lead-gen partnerships with hostels/PGs",
            "Campus activations, social discovery, city-by-city rollout",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    s = add_bullets_slide(
        prs,
        "Roadmap",
        [
            "v1: Backend endpoints + frontend integration demo",
            "v2: Feedback loop, better NLP parsing, trust/safety signals",
            "v3: Payments/reservations, verified profiles, scoring A/B tests",
        ],
    )
    add_accent_bar(s, accent_rgb)
    add_footer(s, footer_text)

    ask = args.ask or "Looking for pilot partners and feedback"
    lines = [
        f"Team: {args.team}" if args.team else "Team: Ahmed Mustafa and Maira Anjum",
        f"The Ask: {ask}",
    ]
    add_title_only_slide(prs, "Team & The Ask", lines)

    contact_lines = [
        args.contact or "Contact: moonahmed786@gmail.com",
        args.website or "https://roommatcher.ai",
        args.number or "+92 332 8371943",
    ]
    contact_lines = [l for l in contact_lines if l]
    add_title_only_slide(prs, "Contact", contact_lines or ["Thank you!"])

    # Output
    out_dir = Path(args.outdir or "docs/pitch").resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    if args.outfile:
        of_path = Path(args.outfile)
        if not of_path.parent or str(of_path.parent) in {".", ""}:
            outfile = out_dir / of_path.name
        else:
            outfile = of_path.resolve()
    else:
        outfile = out_dir / "RoomMatcherAI_PitchDeck.pptx"
    outfile.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(outfile))
    print(f"Wrote deck: {outfile}")


def parse_args():
    p = argparse.ArgumentParser(description="Generate a pitch deck for Room Matcher AI backend")
    p.add_argument("--project-name", default="Room Matcher AI")
    p.add_argument(
        "--tagline",
        default="Minimal multi-agent roommate matcher & room listings search for Pakistan",
    )
    p.add_argument("--author", default="")
    p.add_argument("--team", default="")
    p.add_argument("--ask", default="")
    p.add_argument("--contact", default="")
    p.add_argument("--website", default="")
    p.add_argument("--accent", default="#2F5597")
    p.add_argument("--logo", default="")
    p.add_argument("--footer", default="")
    p.add_argument("--number", default="")
    p.add_argument("--outdir", default="docs/pitch")
    p.add_argument("--outfile", default="")
    return p.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build_deck(args)
